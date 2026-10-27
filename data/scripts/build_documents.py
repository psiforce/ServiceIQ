"""Build the PDF and DOCX corpus documents from their Markdown sources.

Sources live in data/tenants/<tenant>/src/<document_id>.<pdf|docx>.md and are
written to data/tenants/<tenant>/raw/<document_id>.<pdf|docx>. HTML and CSV
documents are authored directly in raw/.

Usage:  .venv/Scripts/python data/scripts/build_documents.py
"""
import html
import re
from pathlib import Path

import docx
import pymupdf
from docx.shared import Pt, RGBColor

DATA = Path(__file__).resolve().parent.parent
CSS = """
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.35; color: #1c2833; }
h1 { font-size: 18pt; color: #1f3a5f; margin-bottom: 4pt; }
h2 { font-size: 13pt; color: #1f3a5f; margin-top: 12pt; }
h3 { font-size: 11pt; color: #1f3a5f; }
table { border-collapse: collapse; width: 100%; }
td, th { border: 0.5pt solid #9aa5b1; padding: 3pt; font-size: 9.5pt; }
th { background-color: #e8eef6; }
p.meta { color: #5d6d7e; font-size: 9pt; }
"""


def parse(md):
    """Yield (kind, payload) blocks from a small Markdown subset."""
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
        elif line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            yield 'h', (level, line[level:].strip())
            i += 1
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            yield 'table', rows
        elif re.match(r'^(- |\d+\. )', line):
            items = []
            while i < len(lines) and re.match(r'^(- |\d+\. )', lines[i]):
                m = re.match(r'^(- |(\d+)\. )(.*)', lines[i])
                items.append((m.group(2), m.group(3)))
                i += 1
            yield 'list', items
        elif line.startswith('_') and line.endswith('_'):
            yield 'meta', line.strip('_')
            i += 1
        else:
            para = []
            while i < len(lines) and lines[i].strip() and not re.match(r'^(#|\||- |\d+\. )', lines[i]):
                para.append(lines[i].strip())
                i += 1
            yield 'p', ' '.join(para)


def inline_html(text):
    return re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', html.escape(text, quote=False))


def to_pdf(md, out):
    parts = []
    for kind, payload in parse(md):
        if kind == 'h':
            parts.append(f'<h{payload[0]}>{inline_html(payload[1])}</h{payload[0]}>')
        elif kind == 'p':
            parts.append(f'<p>{inline_html(payload)}</p>')
        elif kind == 'meta':
            parts.append(f'<p class="meta">{inline_html(payload)}</p>')
        elif kind == 'list':
            tag = 'ol' if payload[0][0] else 'ul'
            parts.append(f'<{tag}>' + ''.join(f'<li>{inline_html(t)}</li>' for _, t in payload) + f'</{tag}>')
        elif kind == 'table':
            head, *body = payload
            parts.append('<table><tr>' + ''.join(f'<th>{inline_html(c)}</th>' for c in head) + '</tr>' +
                         ''.join('<tr>' + ''.join(f'<td>{inline_html(c)}</td>' for c in r) + '</tr>' for r in body) + '</table>')
    story = pymupdf.Story(html='<body>' + ''.join(parts) + '</body>', user_css=CSS)
    writer = pymupdf.DocumentWriter(str(out))
    page = pymupdf.paper_rect('letter')
    where = page + (54, 54, -54, -54)
    more = 1
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()


def add_runs(par, text):
    for i, chunk in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if chunk:
            par.add_run(chunk).bold = bool(i % 2)


def to_docx(md, out):
    d = docx.Document()
    d.styles['Normal'].font.name = 'Calibri'
    d.styles['Normal'].font.size = Pt(11)
    for kind, payload in parse(md):
        if kind == 'h':
            d.add_heading(payload[1], level=min(payload[0] - 1, 3) if payload[0] > 1 else 0)
        elif kind == 'p':
            add_runs(d.add_paragraph(), payload)
        elif kind == 'meta':
            run = d.add_paragraph().add_run(payload)
            run.italic, run.font.size, run.font.color.rgb = True, Pt(9), RGBColor(0x5D, 0x6D, 0x7E)
        elif kind == 'list':
            for num, text in payload:
                add_runs(d.add_paragraph(style='List Number' if num else 'List Bullet'), text)
        elif kind == 'table':
            t = d.add_table(rows=len(payload), cols=len(payload[0]))
            t.style = 'Light Grid Accent 1'
            for r, row in enumerate(payload):
                for c, val in enumerate(row):
                    t.cell(r, c).text = val
    d.save(str(out))


def main():
    count = 0
    for src in sorted(DATA.glob('tenants/*/src/*.md')):
        doc_id, fmt = src.name[:-3].rsplit('.', 1)
        out = src.parent.parent / 'raw' / f'{doc_id}.{fmt}'
        out.parent.mkdir(exist_ok=True)
        (to_pdf if fmt == 'pdf' else to_docx)(src.read_text(encoding='utf-8'), out)
        count += 1
        print(f'built {out.relative_to(DATA)}')
    print(f'{count} documents built')


if __name__ == '__main__':
    main()
