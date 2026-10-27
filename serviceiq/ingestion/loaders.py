"""Load source files into LangChain Documents, keeping document structure.

Each source file becomes one or more Documents ("parts"):
- CSV:  one Document per row, written as "column: value" lines.
- PDF:  one Document per page, converted to Markdown (headings from font size, tables as Markdown tables).
- DOCX: one Document, converted to Markdown (headings from paragraph styles, tables as Markdown tables).
- HTML: one Document holding the raw HTML; boilerplate removal happens in the cleaning step (M2).

Every part carries the manifest metadata plus `part` (row or page number).
"""
import csv
import statistics
import unicodedata
from dataclasses import dataclass
from pathlib import Path

import docx
import pymupdf
from langchain_core.documents import Document

from serviceiq.ingestion.manifest import ManifestEntry, read_manifest, tenant_ids


@dataclass
class LoadedDocument:
    entry: ManifestEntry
    parts: list[Document]

    @property
    def char_count(self) -> int:
        return sum(len(p.page_content) for p in self.parts)


def _markdown_table(rows: list[list[str]]) -> str:
    cells = [[(c or "").replace("\n", " ").strip() for c in row] for row in rows if any(row)]
    if not cells:
        return ""
    lines = ["| " + " | ".join(cells[0]) + " |", "|" + "---|" * len(cells[0])]
    lines += ["| " + " | ".join(row) + " |" for row in cells[1:]]
    return "\n".join(lines)


def _load_csv(path: Path) -> list[tuple[int, str]]:
    with path.open(encoding="utf-8", newline="") as f:
        return [(i, "\n".join(f"{k}: {v}" for k, v in row.items())) for i, row in enumerate(csv.DictReader(f))]


def _pdf_page_markdown(page: pymupdf.Page) -> str:
    tables = page.find_tables().tables
    table_boxes = [pymupdf.Rect(t.bbox) for t in tables]
    # Read each cell from its own box: Table.extract() can drop the spaces between words
    items = [(t.bbox[1], _markdown_table([[page.get_textbox(pymupdf.Rect(c)) if c else "" for c in row.cells]
                                           for row in t.rows])) for t in tables]

    blocks = [b for b in page.get_text("dict")["blocks"] if b["type"] == 0]
    sizes = [s["size"] for b in blocks for line in b["lines"] for s in line["spans"] if s["text"].strip()]
    body_size = statistics.median(sizes) if sizes else 10.0

    for block in blocks:
        if any(pymupdf.Rect(block["bbox"]).intersects(r) for r in table_boxes):
            continue
        spans = [s for line in block["lines"] for s in line["spans"] if s["text"].strip()]
        if not spans:
            continue
        size = max(s["size"] for s in spans)
        all_bold = all(s["flags"] & 16 for s in spans)
        # Join wrapped lines into paragraphs; a bullet character starts a new list item
        paragraphs = []
        for line in block["lines"]:
            line_text = " ".join(s["text"].strip() for s in line["spans"] if s["text"].strip())
            if not line_text:
                continue
            if line_text.startswith(("•", "·")) or not paragraphs:
                paragraphs.append(line_text)
            else:
                paragraphs[-1] += " " + line_text
        for text in paragraphs:
            if size >= body_size * 1.5:
                text = "# " + text
            elif size >= body_size * 1.15:
                text = "## " + text
            elif all_bold and len(text) < 80:
                text = "### " + text
            elif text.startswith(("•", "·")):
                text = "- " + text.lstrip("•· ")
            items.append((block["bbox"][1], text))

    items.sort(key=lambda item: item[0])
    # NFKC turns typographic ligatures such as "ﬁ" into plain letters so search terms match
    return unicodedata.normalize("NFKC", "\n\n".join(text for _, text in items if text))


def _load_pdf(path: Path) -> list[tuple[int, str]]:
    with pymupdf.open(path) as pdf:
        return [(i + 1, _pdf_page_markdown(page)) for i, page in enumerate(pdf)]


def _load_docx(path: Path) -> list[tuple[int, str]]:
    heading_prefix = {"Title": "# ", "Heading 1": "## ", "Heading 2": "### ", "Heading 3": "#### "}
    out = []
    for item in docx.Document(str(path)).iter_inner_content():
        if isinstance(item, docx.table.Table):
            out.append(_markdown_table([[cell.text for cell in row.cells] for row in item.rows]))
            continue
        text = item.text.strip()
        if not text:
            continue
        style = item.style.name if item.style is not None else ""
        if style in heading_prefix:
            text = heading_prefix[style] + text
        elif style.startswith("List Bullet"):
            text = "- " + text
        elif style.startswith("List Number"):
            text = "1. " + text
        out.append(text)
    return [(1, "\n\n".join(out))]


def _load_html(path: Path) -> list[tuple[int, str]]:
    return [(1, path.read_text(encoding="utf-8"))]


LOADERS = {"csv": _load_csv, "pdf": _load_pdf, "docx": _load_docx, "html": _load_html}


def load_document(entry: ManifestEntry) -> LoadedDocument:
    loader = LOADERS.get(entry.format)
    if loader is None:
        raise ValueError(f"{entry.document_id}: unsupported format '{entry.format}'")
    base = entry.metadata()
    parts = [Document(page_content=text, metadata={**base, "part": part}) for part, text in loader(entry.path)]
    return LoadedDocument(entry=entry, parts=parts)


def load_tenant(data_dir: Path, tenant_id: str) -> list[LoadedDocument]:
    return [load_document(entry) for entry in read_manifest(data_dir, tenant_id)]


def load_all(data_dir: Path, tenant: str = "all") -> list[LoadedDocument]:
    tenants = tenant_ids(data_dir) if tenant == "all" else [tenant]
    return [doc for t in tenants for doc in load_tenant(data_dir, t)]
