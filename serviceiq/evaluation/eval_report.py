"""Agent evaluation report: baseline vs improved run on the same golden cases (v2, 50 cases).

Inputs: two saved runs (reports/intake_<stamp>.json from `evaluate-intake`), the golden case files, and a short
analysis file written by the evaluator (failure modes, changes, next steps). Outputs:
  reports/agent_eval_v2.xlsx        case-level results, rubric, judge outputs, human-label sheet, LangSmith links
  docs/ServiceIQ — Agent Evaluation Report (v2).docx   the written report (converted to PDF in Word)
The reply judge (reply_judge.py) grades each run's customer-facing replies; its verdicts are cached next to the run.
"""
import csv
import json
from pathlib import Path

from serviceiq.evaluation import reply_judge
from serviceiq.evaluation.intake_eval import RACE_MESSAGES, TARGETS, AT_MOST, load_case_list
from serviceiq.evaluation.postservice_eval import load_scenarios

METRICS = [  # (key, label, kind) shown in the report, in this order
    ("outcome_accuracy", "Outcome accuracy (dialogues)", "quality"),
    ("postservice_pass_rate", "Post-service scenario pass rate", "quality"),
    ("field_accuracy", "Field accuracy", "quality"),
    ("judge_pass_rate", "Reply quality: LLM-judge pass rate", "quality"),
    ("trajectory_accuracy", "Trajectory accuracy (money and review tools)", "behaviour"),
    ("approval_routing", "Approval routing", "behaviour"),
    ("emergency_exit_rate", "Emergency exit", "safety"),
    ("pii_leaks", "Contact-detail leaks", "safety"),
    ("double_bookings", "Double bookings", "integrity"),
    ("duplicate_charges", "Duplicate charges", "integrity"),
    ("unaudited_money_actions", "Unaudited money actions", "integrity"),
    ("agent_seconds_to_book_p95", "p95 agent time to book (s)", "latency"),
    ("cost_per_case", "Model cost per case ($)", "cost"),
]
EXTRA_TARGETS = {"judge_pass_rate": 0.9, "cost_per_case": 0.01}
EXTRA_AT_MOST = {"cost_per_case"}


def target(key):
    return TARGETS.get(key, EXTRA_TARGETS.get(key))


def at_most(key):
    return key in AT_MOST or key in EXTRA_AT_MOST


def meets(key, value) -> bool | None:
    t = target(key)
    if t is None or value is None:
        return None
    return value <= t if at_most(key) else value >= t


def fmt(key, value) -> str:
    if value is None:
        return "-"
    if key == "cost_per_case":
        return f"${value:.4f}"
    if at_most(key) or key.startswith("agent_seconds"):
        return f"{value:g}"
    return f"{value:.1%}"


# ---------------------------------------------------------------- golden cases
def golden_cases(data_dir: Path, case_ids: set[str]) -> dict[str, dict]:
    """case_id -> {type, tenant, input, expected, expected_behavior} for the cases in the run."""
    cases = {}
    with open(data_dir / "eval" / "intake_dialogues.csv", encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f):
            fields = json.loads(r["expected_fields"] or "{}")
            expected = f"outcome {r['expected_outcome']}" + (f"; fields {fields}" if fields else "") + \
                (f"; must mention {r['must_contain']}" if r["must_contain"] else "")
            cases[r["case_id"]] = {"type": r["scenario_type"], "tenant": r["tenant"],
                                   "input": r["turns"].replace("@", "(when asked) @"), "expected": expected,
                                   "behavior": f"{r['notes']}. Expected outcome: {r['expected_outcome']}."
                                               + (f" Must mention: {r['must_contain']}." if r["must_contain"] else "")}
    for r in load_scenarios(data_dir):
        expected = (f"stage {r['expected_stage']}; invoice {r['expected_invoice'] or 'none'}; tools "
                    f"{r['expected_tools'] or 'none'}; tickets {r['expected_tickets']}"
                    + (f"; reply {r['expected_reply_status']}" if r.get("expected_reply_status") else "")
                    + (f"; must mention {r['must_contain']}" if r["must_contain"] else ""))
        cases[r["case_id"]] = {"type": r["scenario_type"], "tenant": r["tenant"],
                               "input": f"booking: {r['booking']} || {r['events']}", "expected": expected,
                               "behavior": f"{r['notes']}." + (f" The reply should be '{r['expected_reply_status']}'."
                                                               if r.get("expected_reply_status") else "")
                                           + (f" Must mention: {r['must_contain']}." if r["must_contain"] else "")}
    cases["R01"] = {"type": "known_failure", "tenant": "cedar-ridge", "input": " || ".join(RACE_MESSAGES) +
                    " (two customers at once, both pick option 1)",
                    "expected": "one booked, the other re-offered fresh times; no double booking",
                    "behavior": "Race for the same slot."}
    return {k: v for k, v in cases.items() if k in case_ids}


def case_rows(run: dict) -> dict[str, dict]:
    """case_id -> {passed, output, problems, category} from a saved run."""
    rows = {}
    for d in run["dialogues"]:
        problems = ([f"outcome {d['outcome']} (expected {d['expected']})"] if d["outcome"] != d["expected"] else []) \
            + d["field_errors"] + [f"missing phrase: {p}" for p in d["missing_phrases"]]
        last = d["transcript"][-1]["assistant"] if d["transcript"] else ""
        rows[d["case_id"]] = {"passed": d["passed"], "problems": problems,
                              "output": f"{d['outcome']} via {' > '.join(d['statuses'])}; fields "
                                        f"{d['fields_correct']}/{d['fields_expected']}; last reply: {last[:300]}",
                              "category": category(problems), "seconds": d["agent_seconds"], "cost": d["cost_usd"]}
    for p in run["postservice"]:
        problems = p["problems"] + [f"missing phrase: {x}" for x in p["missing_phrases"]]
        rows[p["case_id"]] = {"passed": p["passed"], "problems": problems,
                              "output": f"stage {p['stage']}; invoice {p['invoice'] or 'none'}; tools "
                                        f"{' > '.join(p['tools']) or 'none'}; tickets {p['tickets']}"
                                        + (f"; reply [{(p.get('reply_statuses') or ['-'])[-1]}] {p['replies'][-1][:250]}"
                                           if p["replies"] else ""),
                              "category": category(problems), "seconds": None, "cost": p["cost_usd"]}
    race = run["metrics"].get("race_detail")
    if race:
        rows["R01"] = {"passed": race["handled"], "problems": [] if race["handled"] else ["race not handled"],
                       "output": " > ".join(race["statuses"]), "category": "" if race["handled"] else "integrity",
                       "seconds": None, "cost": 0.0}
    return rows


def category(problems: list[str]) -> str:
    """Failure category from the checks that failed (empty when the case passed)."""
    text = " ".join(problems).lower()
    if not text:
        return ""
    for key, label in (("new service request", "routing after booking"), ("reply:", "routing after booking"),
                       ("outcome", "routing / outcome"), ("tools", "trajectory"), ("invoice", "money"),
                       ("tickets", "hand-off"), ("missing phrase", "incomplete answer"), ("expected", "extraction"),
                       ("stage", "routing / outcome"), ("phone", "privacy")):
        if key in text:
            return label
    return "other"


def summary_metrics(run: dict, judge: list[dict]) -> dict:
    m = dict(run["metrics"])
    cases = len(run["dialogues"]) + len(run["postservice"]) + (1 if "race_detail" in m else 0)
    m["cost_per_case"] = (m.get("total_cost", 0) + m.get("postservice_cost", 0)) / max(1, cases)
    m["judge_pass_rate"] = sum(j["pass"] for j in judge) / len(judge) if judge else None
    m["cases"], m["passed"] = cases, sum(r["passed"] for r in case_rows(run).values())
    return m


def run_judge(run_path: Path, data_dir: Path, settings, cases: dict) -> list[dict]:
    """Grade each reply once; cached in <run>.judge.json."""
    cache = run_path.with_suffix(".judge.json")
    if cache.exists():
        return json.loads(cache.read_text(encoding="utf-8"))
    from serviceiq.llm import openai_client
    client, run, out = openai_client(settings), json.loads(run_path.read_text(encoding="utf-8")), []
    for item in reply_judge.reply_items(run):
        case = cases[item["case_id"]]
        verdict, usage = reply_judge.judge_reply(client, settings.judge_model, item, case["behavior"],
                                                 reply_judge.company_facts(data_dir, item["tenant"]), case["input"])
        out.append(item | verdict | {"judge_cost": usage.cost or 0.0})
    cache.write_text(json.dumps(out, indent=1), encoding="utf-8")
    return out


# ---------------------------------------------------------------- spreadsheet
def write_workbook(path: Path, cases: dict, base: dict, improved: dict, base_m: dict, imp_m: dict,
                   judges: tuple[list, list], analysis: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head, wrap = Font(bold=True, color="FFFFFF"), Alignment(wrap_text=True, vertical="top")
    fill = PatternFill("solid", fgColor="1F3A5F")
    green, red = PatternFill("solid", fgColor="E7F6EE"), PatternFill("solid", fgColor="FDECEC")

    def sheet(wb, title, headers, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill, cell.alignment = head, fill, wrap
        for row in rows:
            ws.append(row)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
                if cell.value in ("PASS", "FAIL"):
                    cell.fill = green if cell.value == "PASS" else red
        ws.freeze_panes = "B2"
        return ws

    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "Summary", ["Metric", "Kind", "Pass bar", "Baseline", "Improved", "Delta", "Baseline meets", "Improved meets"],
          [[label, kind, (("≤ " if at_most(k) else "≥ ") + fmt(k, target(k))) if target(k) is not None else "-",
            fmt(k, base_m.get(k)), fmt(k, imp_m.get(k)), delta(k, base_m.get(k), imp_m.get(k)),
            yes(meets(k, base_m.get(k))), yes(meets(k, imp_m.get(k)))] for k, label, kind in METRICS]
          + [["Cases passed", "overall", "-", f"{base_m['passed']}/{base_m['cases']}",
              f"{imp_m['passed']}/{imp_m['cases']}", f"{imp_m['passed'] - base_m['passed']:+d}", "", ""]],
          [44, 12, 12, 14, 14, 12, 14, 14])
    order = sorted(cases, key=lambda c: (["happy", "edge", "known_failure", "adversarial"].index(cases[c]["type"]), c))
    sheet(wb, "Cases", ["Case", "Type", "Tenant", "Input", "Expected output / behaviour", "Baseline output",
                        "Baseline", "Failure category", "Baseline problems", "Improved output", "Improved",
                        "Improved problems", "Change"],
          [[c, cases[c]["type"], cases[c]["tenant"], cases[c]["input"], cases[c]["expected"],
            base.get(c, {}).get("output", "not run"), pf(base.get(c)), base.get(c, {}).get("category", ""),
            "; ".join(base.get(c, {}).get("problems", [])), improved.get(c, {}).get("output", "not run"),
            pf(improved.get(c)), "; ".join(improved.get(c, {}).get("problems", [])),
            change(base.get(c), improved.get(c))] for c in order],
          [7, 13, 13, 50, 40, 55, 9, 18, 35, 55, 9, 35, 12])
    sheet(wb, "Rubric", ["Metric / criterion", "Scoring method", "Rule", "Pass bar"], analysis["rubric"],
          [34, 22, 80, 14])
    for label, judged in (("Judge baseline", judges[0]), ("Judge improved", judges[1])):
        sheet(wb, label, ["Case", "Type", "Reply status", "Customer", "Reply", "Grounded", "Right action", "Safe",
                          "Clear", "Overall (1-5)", "Judge", "Reason"],
              [[j["case_id"], j["scenario_type"], j["status"], j["customer"], j["reply"], j["grounded"],
                j["right_action"], j["safe"], j["clear"], j["overall"], "PASS" if j["pass"] else "FAIL", j["reason"]]
               for j in judged], [7, 13, 14, 35, 70, 10, 10, 8, 8, 10, 8, 50])
    sample = judges[1][:30]
    ws = sheet(wb, "Human labels", ["Case", "Reply", "Judge overall", "Judge pass", "Human overall (1-5)",
                                    "Human pass (PASS/FAIL)", "Human notes"],
               [[j["case_id"], j["reply"], j["overall"], "PASS" if j["pass"] else "FAIL", None, None, None]
                for j in sample], [7, 80, 12, 11, 16, 18, 40])
    ws.cell(row=len(sample) + 3, column=1, value="Fill columns E and F using the Rubric tab, then run "
                                                  "`python -m serviceiq eval-report ... --labels <this file>` to "
                                                  "compute judge-human agreement.")
    sheet(wb, "Judge prompt", ["Judge system prompt (model: settings.judge_model)", "Output schema"],
          [[reply_judge.SYSTEM, json.dumps(reply_judge.SCHEMA, indent=1)]], [100, 60])
    sheet(wb, "Changes", ["#", "Change", "Lever", "Files", "Targets", "Predicted", "Measured"],
          [[c["id"], c["change"], c["lever"], c["files"], c["targets"], c["predicted"], c.get("measured", "")]
           for c in analysis["changes"]], [5, 50, 14, 30, 22, 30, 30])
    sheet(wb, "LangSmith", ["Item", "Value"], [[k, v] for k, v in analysis["langsmith"].items()], [30, 120])
    path.parent.mkdir(exist_ok=True)
    wb.save(path)


def pf(row) -> str:
    return "" if row is None else "PASS" if row["passed"] else "FAIL"


def yes(flag) -> str:
    return "" if flag is None else "yes" if flag else "NO"


def change(before, after) -> str:
    if before is None or after is None:
        return ""
    return {(False, True): "fixed", (True, False): "REGRESSED", (False, False): "still failing"}.get(
        (before["passed"], after["passed"]), "")


def delta(key, a, b) -> str:
    if a is None or b is None:
        return "-"
    d = b - a
    if key == "cost_per_case":
        return f"{d:+.4f}"
    if at_most(key) or key.startswith("agent_seconds"):
        return f"{d:+g}"
    return f"{d * 100:+.1f} pts"


# ---------------------------------------------------------------- judge vs human labels
def agreement(labels_path: Path, judged: list[dict]) -> dict | None:
    """Judge-human agreement from the filled 'Human labels' sheet (None until a person has labelled)."""
    from openpyxl import load_workbook
    ws = load_workbook(labels_path, read_only=True)["Human labels"]
    by_case = {j["case_id"]: j for j in judged}
    pairs = [(by_case[r[0]], r[4], str(r[5]).strip().upper()) for r in ws.iter_rows(min_row=2, values_only=True)
             if r[0] in by_case and r[5] not in (None, "")]
    if not pairs:
        return None
    same = sum((j["pass"] and human == "PASS") or (not j["pass"] and human == "FAIL") for j, _, human in pairs)
    scored = [(j["overall"], int(h)) for j, h, _ in pairs if h not in (None, "")]
    return {"labelled": len(pairs), "pass_agreement": same / len(pairs),
            "judge_pass_human_fail": sum(j["pass"] and h == "FAIL" for j, _, h in pairs),
            "judge_fail_human_pass": sum(not j["pass"] and h == "PASS" for j, _, h in pairs),
            "mean_abs_score_diff": sum(abs(a - b) for a, b in scored) / len(scored) if scored else None}


# ---------------------------------------------------------------- written report (python-docx)
def write_docx(path: Path, analysis: dict, cases: dict, base: dict, improved: dict, base_m: dict, imp_m: dict,
               judges: tuple[list, list], agree: dict | None) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10.5)
    navy = RGBColor(0x1F, 0x3A, 0x5F)

    def h(text, level=1):
        doc.add_heading(text, level).runs[0].font.color.rgb = navy

    def p(text, bold_lead=None):
        para = doc.add_paragraph()
        if bold_lead:
            para.add_run(bold_lead + " ").bold = True
        para.add_run(text)
        return para

    def bullets(items):
        for item in items:
            doc.add_paragraph(item, style="List Bullet")

    def table(headers, rows, size=8.5):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style, t.alignment = "Light Grid Accent 1", WD_TABLE_ALIGNMENT.CENTER
        for cell, text in zip(t.rows[0].cells, headers):
            cell.text = text
        for row in rows:
            for cell, text in zip(t.add_row().cells, row):
                cell.text = str(text)
        for row in t.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(size)
        doc.add_paragraph()

    a = analysis
    doc.add_heading("ServiceIQ Agent Evaluation Report", 0)
    p(f"{a['date']} · golden dataset v2 (50 cases) · baseline and improved runs on identical cases · "
      f"LangSmith dataset {a['langsmith']['Dataset']}")

    h("1. Evaluation report")
    table(["Item", "Detail"], [["Use case", a["use_case"]], ["Agent under test", a["agent"]],
                               ["Target user", a["target_user"]], ["User outcome", a["outcome"]]])
    p(a["one_liner"], "Evaluation one-liner.")
    h("Metrics, scoring and pass bars", 2)
    table(["Metric", "Kind", "Scoring method", "Pass bar"],
          [[label, kind, a["methods"].get(k, ""), (("≤ " if at_most(k) else "≥ ") + fmt(k, target(k)))
            if target(k) is not None else "-"] for k, label, kind in METRICS])
    h("Rubric", 2)
    table(["Metric / criterion", "Method", "Rule", "Pass bar"], a["rubric"], size=8)

    h("2. Golden dataset and scored results")
    p(a["dataset"])
    counts = {}
    for c in cases.values():
        counts[c["type"]] = counts.get(c["type"], 0) + 1
    table(["Scenario type", "Cases", "Baseline passed", "Improved passed"],
          [[t, counts.get(t, 0), sum(base[c]["passed"] for c in cases if cases[c]["type"] == t and c in base),
            sum(improved[c]["passed"] for c in cases if cases[c]["type"] == t and c in improved)]
           for t in ("happy", "edge", "known_failure", "adversarial")]
          + [["total", len(cases), base_m["passed"], imp_m["passed"]]])
    table(["Metric", "Baseline", "Improved", "Delta", "Pass bar met (improved)"],
          [[label, fmt(k, base_m.get(k)), fmt(k, imp_m.get(k)), delta(k, base_m.get(k), imp_m.get(k)),
            yes(meets(k, imp_m.get(k)))] for k, label, _ in METRICS])
    h("Case-level results", 2)
    p("Full inputs, expected behaviour and outputs for every case are in the spreadsheet (tab Cases) and in the "
      "LangSmith experiments.")
    order = sorted(cases, key=lambda c: (["happy", "edge", "known_failure", "adversarial"].index(cases[c]["type"]), c))
    table(["Case", "Type", "Expected", "Baseline", "Failure category", "Improved", "Change"],
          [[c, cases[c]["type"], cases[c]["expected"][:90], pf(base.get(c)), base.get(c, {}).get("category", ""),
            pf(improved.get(c)), change(base.get(c), improved.get(c))] for c in order], size=7.5)

    h("3. Failure analysis and measured improvements")
    h("Top failure modes (baseline)", 2)
    for i, mode in enumerate(a["failure_modes"], 1):
        p(f"{mode['description']} Cases: {', '.join(mode['cases'])}. Example ({mode['example_case']}): "
          f"{mode['example']} Rough cost: {mode['cost']}", f"{i}. {mode['name']}.")
    h("Targeted changes", 2)
    table(["#", "Change", "Lever", "Files", "Targets", "Predicted", "Measured"],
          [[c["id"], c["change"], c["lever"], c["files"], c["targets"], c["predicted"], c.get("measured", "")]
           for c in a["changes"]], size=8)
    h("Before / after", 2)
    fixed = [c for c in order if change(base.get(c), improved.get(c)) == "fixed"]
    regressed = [c for c in order if change(base.get(c), improved.get(c)) == "REGRESSED"]
    still = [c for c in order if change(base.get(c), improved.get(c)) == "still failing"]
    bullets([f"Cases passed: {base_m['passed']}/{base_m['cases']} → {imp_m['passed']}/{imp_m['cases']}.",
             f"Fixed: {', '.join(fixed) or 'none'}.", f"Regressions: {', '.join(regressed) or 'none'}.",
             f"Still failing: {', '.join(still) or 'none'}."] + a.get("before_after_notes", []))
    h("Remaining failures and next steps", 2)
    bullets(a["next_steps"])

    h("4. Supporting materials")
    bullets(a["materials"])
    h("LLM judge: prompt, outputs and human comparison", 2)
    p(reply_judge.SYSTEM)
    jb, ji = judges
    p(f"Judge model {a['judge_model']}. Baseline: {sum(j['pass'] for j in jb)}/{len(jb)} replies pass; improved: "
      f"{sum(j['pass'] for j in ji)}/{len(ji)}. Every verdict with its reason is in the spreadsheet (tabs Judge "
      "baseline / Judge improved).")
    if agree:
        p(f"{agree['labelled']} replies labelled by a person: pass/fail agreement {agree['pass_agreement']:.0%}; judge "
          f"passed {agree['judge_pass_human_fail']} the person failed and failed {agree['judge_fail_human_pass']} the "
          f"person passed; mean score difference "
          + (f"{agree['mean_abs_score_diff']:.2f}." if agree["mean_abs_score_diff"] is not None else "-."),
          "Judge vs human.")
    else:
        p("Pending: 30 improved-run replies are in the spreadsheet's Human labels tab for a person to grade with the "
          "same rubric. Until then the judge's scores are reported but not relied on for any pass/fail decision.",
          "Judge vs human.")
    failing = [j for j in ji if not j["pass"]]
    if failing:
        table(["Case", "Overall", "Reason"], [[j["case_id"], j["overall"], j["reason"]] for j in failing], size=8)

    h("5. LangSmith evidence")
    table(["Item", "Value"], [[k, v] for k, v in a["langsmith"].items()], size=8)
    p(a.get("langsmith_note", ""))
    path.parent.mkdir(exist_ok=True)
    doc.save(path)
